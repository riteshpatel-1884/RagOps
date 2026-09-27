"""
API layer — exposes retrieval eval, generation eval, and experiment sweeps
as HTTP endpoints for the Next.js frontend.

Design choice: rather than rearchitecting the CLI scripts, uploaded
documents/test-datasets are written straight into data/documents.json and
data/test_dataset.json — the exact files evaluate.py, evaluate_generation.py,
and experiment.py already read via their own load_data(). That means this
API layer is a thin wrapper: it doesn't duplicate any evaluation, retrieval,
or diagnostic logic, it just calls the same functions the CLI does and
returns their results as JSON instead of printing them.

Run with:
    uvicorn api:app --reload --port 8000

CORS is wide open (allow_origins=["*"]) because this is a local dev tool,
not a deployed multi-tenant service — tighten this before deploying anywhere
real.
"""
import io
import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import List, Dict

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

from app.pipeline import PipelineConfig, RetrievalPipeline
import app.evaluate as evaluate_mod
import app.evaluate_generation as evaluate_generation_mod
import app.experiment as experiment_mod

DATA_DIR = Path(__file__).parent / "data"
RESULTS_DIR = Path(__file__).parent / "experiment_results"
DATA_DIR.mkdir(parents=True, exist_ok=True)
RESULTS_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="RAG Evaluation & Optimization Engine API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------------------------------------------------------------------------
# Request/response models
# ---------------------------------------------------------------------------

class TestDatasetItem(BaseModel):
    id: str
    question: str
    relevant_doc_ids: List[str]


class RetrievalEvalRequest(BaseModel):
    chunk_size: int = 400
    chunk_overlap: int = 80
    embedder_name: str = "bge-small-en-v1.5"
    top_k: int = 5


class GenerationEvalRequest(RetrievalEvalRequest):
    generator_name: str = "groq"
    judge: str = "offline"


class AskRequest(BaseModel):
    question: str
    chunk_size: int = 400
    chunk_overlap: int = 80
    embedder_name: str = "bge-small-en-v1.5"
    top_k: int = 5
    generator_name: str = "groq"


class ExperimentRequest(BaseModel):
    chunk_sizes: List[int] = [200, 400]
    chunk_overlaps: List[int] = [40, 80]
    embedders: List[str] = ["bge-small-en-v1.5"]
    top_ks: List[int] = [3, 5]
    generators: List[str] = ["groq"]
    judge: str = "offline"
    skip_generation: bool = False
    sort_by: str = "recall"


# ---------------------------------------------------------------------------
# Config options (drives the UI's dropdowns from one source of truth)
# ---------------------------------------------------------------------------

@app.get("/api/config-options")
def config_options():
    return {
        "embedders": [
            {"value": "bge-small-en-v1.5", "label": "BGE Small EN v1.5 (fastest, lightest)"},
            {"value": "bge-m3", "label": "BGE M3 (strongest, multilingual)"},
            {"value": "qwen3-embedding-0.6b", "label": "Qwen3 Embedding 0.6B"},
        ],
        "generators": [
            {"value": "groq", "label": "Groq (requires GROQ_API_KEY)"},
        ],
        "judges": [
            {"value": "offline", "label": "Offline (lexical overlap + embedding similarity)"},
            {"value": "llm", "label": "LLM judge (reuses the generator's model)"},
        ],
    }


# ---------------------------------------------------------------------------
# Documents — text extraction for any uploaded file format
# ---------------------------------------------------------------------------
#
# Every format below is reduced to plain text before it goes into
# documents.json, so the rest of the pipeline (chunking.py, embedder.py,
# etc.) never needs to know or care what format a document originally was.

def _decode_bytes(raw: bytes, filename: str) -> str:
    """Best-effort plain-text decode, trying the most common encodings in order."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HTTPException(400, f"Could not decode '{filename}' as text.")


def _extract_pdf_text(raw: bytes, filename: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise HTTPException(
            500, "PDF support requires the 'pypdf' package on the server. Install it with: pip install pypdf"
        )
    try:
        reader = PdfReader(io.BytesIO(raw))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as e:
        raise HTTPException(400, f"Could not read '{filename}' as a PDF: {e}")

    text = "\n\n".join(p.strip() for p in pages if p.strip())
    if not text.strip():
        raise HTTPException(400, f"No extractable text found in '{filename}' (it may be a scanned/image-only PDF).")
    return text


def _extract_docx_text(raw: bytes, filename: str) -> str:
    try:
        import docx
    except ImportError:
        raise HTTPException(
            500, "DOCX support requires the 'python-docx' package on the server. Install it with: pip install python-docx"
        )
    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as e:
        raise HTTPException(400, f"Could not read '{filename}' as a Word document: {e}")

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise HTTPException(400, f"No extractable text found in '{filename}'.")
    return text


def _extract_html_text(raw: bytes, filename: str) -> str:
    from html.parser import HTMLParser

    class _TextExtractor(HTMLParser):
        def __init__(self):
            super().__init__()
            self.chunks: List[str] = []
            self._skip = False

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style"):
                self._skip = True

        def handle_endtag(self, tag):
            if tag in ("script", "style"):
                self._skip = False

        def handle_data(self, data):
            if not self._skip and data.strip():
                self.chunks.append(data.strip())

    parser = _TextExtractor()
    parser.feed(_decode_bytes(raw, filename))
    text = "\n".join(parser.chunks)
    if not text.strip():
        raise HTTPException(400, f"No extractable text found in '{filename}'.")
    return text


def _extract_json_text(raw: bytes, filename: str) -> str:
    """Flatten arbitrary JSON into readable 'path: value' lines. Falls back to
    raw text if the file isn't actually valid JSON despite its extension."""
    raw_str = _decode_bytes(raw, filename)
    try:
        data = json.loads(raw_str)
    except json.JSONDecodeError:
        return raw_str

    def _flatten(value, path=""):
        if isinstance(value, dict):
            for k, v in value.items():
                yield from _flatten(v, f"{path}.{k}" if path else str(k))
        elif isinstance(value, list):
            for i, v in enumerate(value):
                yield from _flatten(v, f"{path}[{i}]")
        elif value not in (None, ""):
            yield f"{path}: {value}" if path else str(value)

    lines = list(_flatten(data))
    return "\n".join(lines) if lines else raw_str


# Extensions that need real parsing. Anything not listed here (.txt, .md,
# .csv, .tsv, and any unrecognized extension) is treated as plain text.
_EXTRACTORS = {
    ".pdf": _extract_pdf_text,
    ".docx": _extract_docx_text,
    ".html": _extract_html_text,
    ".htm": _extract_html_text,
    ".json": _extract_json_text,
}


def extract_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    extractor = _EXTRACTORS.get(suffix)
    if extractor:
        return extractor(raw, filename)
    return _decode_bytes(raw, filename)


# ---------------------------------------------------------------------------
# Documents
# ---------------------------------------------------------------------------

@app.post("/api/documents/upload")
async def upload_documents(files: List[UploadFile] = File(...)):
    """
    Replaces the current corpus with the uploaded files. Each file's name
    (without extension) becomes its doc_id. Text is extracted based on file
    type: .pdf and .docx have their text pulled out, .html/.htm is stripped
    of markup, .json is flattened into readable lines, and everything else
    (.txt, .md, .csv, or any other extension) is read as plain text.
    """
    documents: Dict[str, str] = {}
    for f in files:
        raw = await f.read()
        if not raw:
            raise HTTPException(400, f"'{f.filename}' is empty.")
        documents[Path(f.filename).stem] = extract_text(f.filename, raw)

    if not documents:
        raise HTTPException(400, "No files received.")

    with open(DATA_DIR / "documents.json", "w") as out:
        json.dump(documents, out, indent=2)

    return {"count": len(documents), "doc_ids": list(documents.keys())}


@app.get("/api/documents")
def get_documents():
    path = DATA_DIR / "documents.json"
    if not path.exists():
        return {}
    with open(path) as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# Test dataset
# ---------------------------------------------------------------------------

@app.post("/api/test-dataset")
def save_test_dataset(items: List[TestDatasetItem]):
    """Replaces the current test dataset with the given questions."""
    if not items:
        raise HTTPException(400, "Test dataset can't be empty.")
    with open(DATA_DIR / "test_dataset.json", "w") as f:
        json.dump([item.model_dump() for item in items], f, indent=2)
    return {"count": len(items)}


@app.post("/api/test-dataset/upload")
async def upload_test_dataset(file: UploadFile = File(...)):
    """Alternative to /api/test-dataset: upload an existing test_dataset.json file directly."""
    raw = await file.read()
    try:
        items = json.loads(raw)
    except json.JSONDecodeError:
        raise HTTPException(400, "File isn't valid JSON.")

    required = {"id", "question", "relevant_doc_ids"}
    for item in items:
        if not required.issubset(item.keys()):
            raise HTTPException(400, f"Each item needs {required}. Got: {list(item.keys())}")

    with open(DATA_DIR / "test_dataset.json", "w") as f:
        json.dump(items, f, indent=2)
    return {"count": len(items)}


@app.get("/api/test-dataset")
def get_test_dataset():
    path = DATA_DIR / "test_dataset.json"
    if not path.exists():
        return []
    with open(path) as f:
        return json.load(f)


# ---------------------------------------------------------------------------
# Phase 1 + 2: single-config evaluation
# ---------------------------------------------------------------------------

def _require_data():
    if not (DATA_DIR / "documents.json").exists():
        raise HTTPException(400, "No documents uploaded yet. POST /api/documents/upload first.")
    if not (DATA_DIR / "test_dataset.json").exists():
        raise HTTPException(400, "No test dataset uploaded yet. POST /api/test-dataset first.")


@app.post("/api/evaluate/retrieval")
def evaluate_retrieval(req: RetrievalEvalRequest):
    _require_data()
    config = PipelineConfig(
        chunk_size=req.chunk_size, chunk_overlap=req.chunk_overlap,
        embedder_name=req.embedder_name, top_k=req.top_k,
    )
    try:
        overall, per_query = evaluate_mod.run_evaluation(config, verbose=False)
    except Exception as e:
        raise HTTPException(500, str(e))
    return {"overall": overall, "per_query": per_query}


@app.post("/api/evaluate/generation")
def evaluate_generation(req: GenerationEvalRequest):
    _require_data()
    config = PipelineConfig(
        chunk_size=req.chunk_size, chunk_overlap=req.chunk_overlap,
        embedder_name=req.embedder_name, top_k=req.top_k, generator_name=req.generator_name,
    )
    try:
        result = evaluate_generation_mod.run_generation_evaluation(config, judge=req.judge, verbose=False)
    except Exception as e:
        raise HTTPException(500, str(e))
    return result


@app.post("/api/ask")
def ask(req: AskRequest):
    """Ad-hoc single-question playground — no test dataset required."""
    if not (DATA_DIR / "documents.json").exists():
        raise HTTPException(400, "No documents uploaded yet. POST /api/documents/upload first.")

    with open(DATA_DIR / "documents.json") as f:
        documents = json.load(f)

    config = PipelineConfig(
        chunk_size=req.chunk_size, chunk_overlap=req.chunk_overlap,
        embedder_name=req.embedder_name, top_k=req.top_k, generator_name=req.generator_name,
    )
    try:
        pipeline = RetrievalPipeline(config)
        pipeline.index(documents)
        result = pipeline.answer(req.question)
    except Exception as e:
        raise HTTPException(500, str(e))

    return {
        "question": result["question"],
        "answer": result["answer"],
        "retrieved_hits": result["retrieved_hits"],
    }


# ---------------------------------------------------------------------------
# Phase 3: experiment sweeps
# ---------------------------------------------------------------------------

@app.post("/api/experiment")
def run_experiment(req: ExperimentRequest):
    _require_data()
    documents, test_set = evaluate_mod.load_data()

    configs = []
    import itertools
    for cs, co, emb, k, gen in itertools.product(req.chunk_sizes, req.chunk_overlaps, req.embedders, req.top_ks, req.generators):
        if co >= cs:
            continue
        configs.append(PipelineConfig(chunk_size=cs, chunk_overlap=co, embedder_name=emb, top_k=k, generator_name=gen))

    if not configs:
        raise HTTPException(400, "No valid configs — every chunk_overlap must be smaller than every chunk_size.")

    all_results = experiment_mod.run_sweep(configs, documents, test_set, req.judge, req.skip_generation)
    results = [r for r in all_results if "error" not in r]
    failed = [r for r in all_results if "error" in r]

    if not results:
        raise HTTPException(500, f"All {len(configs)} configs failed. First error: {failed[0]['error'] if failed else 'unknown'}")

    sort_by = req.sort_by if req.sort_by in results[0] else "recall"
    ranked = sorted(results, key=lambda r: r.get(sort_by, 0), reverse=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"experiment_{timestamp}.json"
    with open(RESULTS_DIR / filename, "w") as f:
        json.dump(ranked, f, indent=2)

    return {"results": ranked, "failed": failed, "sort_by": sort_by, "filename": filename}


@app.get("/api/experiments")
def list_experiments():
    files = sorted(RESULTS_DIR.glob("experiment_*.json"), reverse=True)
    return [{"filename": f.name, "modified": datetime.fromtimestamp(f.stat().st_mtime).isoformat()} for f in files]


@app.get("/api/experiments/{filename}")
def get_experiment(filename: str):
    path = RESULTS_DIR / filename
    if not path.exists() or path.parent != RESULTS_DIR:
        raise HTTPException(404, "Experiment file not found.")
    with open(path) as f:
        return json.load(f)


@app.get("/api/health")
def health():
    return {"status": "ok"}